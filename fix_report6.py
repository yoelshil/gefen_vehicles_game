from pptx import Presentation
from pptx.oxml.ns import qn
from copy import deepcopy

INPUT = "/root/.claude/uploads/fc0f1583-bebd-475e-866a-e575167d3e15/93378ada-report_6.pptx"
OUTPUT = "/home/user/gefen_vehicles_game/report_7.pptx"

prs = Presentation(INPUT)
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def find_table(slide):
    for s in slide.shapes:
        if s.has_table:
            return s.table
    raise ValueError("No table found")


def set_cell(table, row_idx, col_idx, text):
    cell = table.cell(row_idx, col_idx)
    tf = cell.text_frame
    para = tf.paragraphs[0]
    if para.runs:
        para.runs[0].text = text
        for r in list(para.runs[1:]):
            para._p.remove(r._r)
    else:
        para.text = text


def clone_last_row(table, values):
    tbl = table._tbl
    tr_list = tbl.findall(qn("a:tr"))
    new_tr = deepcopy(tr_list[-1])
    tc_list = new_tr.findall(qn("a:tc"))
    for tc, val in zip(tc_list, values):
        for t_elem in tc.iter(qn("a:t")):
            t_elem.text = val
            break
    tbl.append(new_tr)


def add_paragraph_to_shape(shape, text, clone_from_idx=-1):
    txBody = shape.text_frame._txBody
    paras = txBody.findall(qn("a:p"))
    ref = paras[clone_from_idx]
    new_p = deepcopy(ref)
    for t in new_p.iter(qn("a:t")):
        t.text = text
        break
    ref.addnext(new_p)
    return new_p


# ── SLIDE 3 (index 2): Research Questions ──

slide3 = prs.slides[2]
tf3 = None
for s in slide3.shapes:
    if s.has_text_frame and s.name == "Content Placeholder 4":
        tf3 = s.text_frame
        break

txBody = tf3._txBody
all_paras = txBody.findall(qn("a:p"))

# Para 8 is the red line to move. Identify it by content.
move_para = None
move_idx = None
for i, p in enumerate(all_paras):
    full = "".join(t.text for t in p.iter(qn("a:t")))
    if "first exposure improve estimates relative to the uninformed blind prior" in full:
        move_para = p
        move_idx = i
        break

# Find "Learning" header para (para 7)
learning_header = None
learning_idx = None
for i, p in enumerate(all_paras):
    full = "".join(t.text for t in p.iter(qn("a:t")))
    if full.strip() == "Learning":
        learning_header = p
        learning_idx = i
        break

# Step 1: Move para 8 to just before the Learning header, append "(Blind vs T1)"
txBody.remove(move_para)
learning_header.addprevious(move_para)
# Append " (Blind vs T1)" to run text
for t in move_para.iter(qn("a:t")):
    t.text = t.text.rstrip() + " (Blind vs T1)"
    break
# Remove red color from the moved paragraph's run
for rPr in move_para.iter(qn("a:rPr")):
    for fill in rPr.findall(qn("a:solidFill")):
        rPr.remove(fill)

# Re-read paragraphs after the move
all_paras = txBody.findall(qn("a:p"))

# Step 2 & 3: Remove " No" and " Yes" from learning questions
for p in all_paras:
    runs = p.findall(qn("a:r"))
    if len(runs) >= 3:
        texts = ["".join(t.text for t in r.iter(qn("a:t"))) for r in runs]
        if texts[-1].strip() in ("No", "Yes"):
            # Remove the last two runs (space + answer)
            for r in runs[-2:]:
                p.remove(r)

# Step 4: Add new Learning question at end (before the last para about monotonic trend)
all_paras = txBody.findall(qn("a:p"))
# Find the monotonic trend paragraph
monotonic_para = None
for p in all_paras:
    full = "".join(t.text for t in p.iter(qn("a:t")))
    if "monotonic trend" in full:
        monotonic_para = p
        break

# Clone a learning question paragraph for formatting
ref_para = None
for p in all_paras:
    full = "".join(t.text for t in p.iter(qn("a:t")))
    if "second exposure with feedback" in full:
        ref_para = p
        break

new_q = deepcopy(ref_para)
for t in new_q.iter(qn("a:t")):
    t.text = "Does the two-exposure sequence improve estimates relative to the icon prior? (Icon vs T2)"
    break
# Remove extra runs if any
runs = new_q.findall(qn("a:r"))
for r in runs[1:]:
    new_q.remove(r)

monotonic_para.addprevious(new_q)

print("✓ Slide 3: Research Questions updated")

# ── SLIDE 4 (index 3): Prior Table ──

tbl4 = find_table(prs.slides[3])
# Row indices: 0=header, 1=P1, 2=P2, 3=P3, 4=P4, 5=P5, 6=P6
set_cell(tbl4, 1, 1, "What distribution do participants assign before seeing any visual information?")
set_cell(tbl4, 2, 1, "Is the blind prior more uniform than expected by chance?")
set_cell(tbl4, 3, 8, "#9")
set_cell(tbl4, 4, 1, "What distributional shapes do participants' priors take?")
set_cell(tbl4, 4, 8, "#9")
set_cell(tbl4, 5, 1, "Is the two-type dominance (Uniform + Unimodal-mild) surprising?")
set_cell(tbl4, 5, 8, "#9")
set_cell(tbl4, 6, 1, "Is each category's observed proportion surprising? (Unimodal-mild)")
set_cell(tbl4, 6, 8, "#9")

print("✓ Slide 4: Prior table updated")

# ── SLIDE 5 (index 4): Icon Prior Table ──

tbl5 = find_table(prs.slides[4])
# Row indices: 0=header, 1=P7, 2=P2a, 3=P2b
set_cell(tbl5, 1, 1, "Does seeing the item shapes change the prior distribution?")
set_cell(tbl5, 1, 8, "#10")
set_cell(tbl5, 2, 1, "Is the icon prior also more uniform than expected by chance?")
set_cell(tbl5, 2, 8, "#10")
set_cell(tbl5, 3, 1, "Is the blind-to-icon reduction significant?")
set_cell(tbl5, 3, 8, "#10")

clone_last_row(tbl5, [
    "P2c",
    "Does first exposure improve over uninformed blind prior? (Blind vs T1)",
    "Paired perm (two-sided)",
    "3.1",
    "Fill from v25",
    "Fill",
    "Fill Cliff's d",
    "Fill",
    "#11",
])

print("✓ Slide 5: Icon Prior table updated + P2c row added")

# ── SLIDE 6 (index 5): Learning Table ──

tbl6 = find_table(prs.slides[5])
# Row indices: 0=header, 1=L1, 2=L2, 3=L3, 4=S1, 5=S2
set_cell(tbl6, 1, 1, "Does the first exposure shift estimates relative to the icon prior?")
set_cell(tbl6, 1, 8, "#11")
set_cell(tbl6, 2, 1, "Is any observed change in the direction of [1,1,4,6]?")
set_cell(tbl6, 2, 8, "#11")
set_cell(tbl6, 3, 1, "Does a second exposure with feedback improve estimates beyond T1?")
set_cell(tbl6, 3, 8, "#12")
set_cell(tbl6, 4, 1, "What is the overall learning from blind to final estimate?")
set_cell(tbl6, 4, 8, "#12")
set_cell(tbl6, 5, 8, "#12")

clone_last_row(tbl6, [
    "S3",
    "Does two-exposure sequence improve over icon prior? (Icon vs T2)",
    "Paired perm + Cliff's d",
    "3.2b",
    "Fill from v25",
    "Fill",
    "Fill",
    "Fill",
    "#12",
])

print("✓ Slide 6: Learning table updated + S3 row added")

# ── SLIDE 10 (index 9): Icon Prior content ──

slide10 = prs.slides[9]
for s in slide10.shapes:
    if s.has_text_frame and "Icon prior SAD" in s.text:
        txBody10 = s.text_frame._txBody
        paras10 = txBody10.findall(qn("a:p"))
        ref10 = paras10[-1]
        new_p10 = deepcopy(ref10)
        for t in new_p10.iter(qn("a:t")):
            t.text = "Blind to T1 (combined effect): Fill diff, Fill p, Fill Cliff's d from v25."
            break
        runs10 = new_p10.findall(qn("a:r"))
        for r in runs10[1:]:
            new_p10.remove(r)
        txBody10.append(new_p10)
        break

print("✓ Slide 10: P2c stat line added")

# ── SLIDE 12 (index 11): Learning Overall ──

slide12 = prs.slides[11]

# Add Icon-to-T2 line to TextBox 2 (the results box)
for s in slide12.shapes:
    if s.has_text_frame and "T1 to T2" in s.text:
        txBody12 = s.text_frame._txBody
        paras12 = txBody12.findall(qn("a:p"))
        # Find the "Direction:" paragraph as reference (has actual text)
        ref12 = None
        for p in paras12:
            full = "".join(t.text for t in p.iter(qn("a:t")))
            if "Direction:" in full:
                ref12 = p
                break
        if ref12 is None:
            ref12 = paras12[0]
        new_p12 = deepcopy(ref12)
        # Set text in first run, remove extras
        runs12 = new_p12.findall(qn("a:r"))
        if runs12:
            for t in runs12[0].iter(qn("a:t")):
                t.text = "Icon to T2 (exposure learning): Fill diff, Fill p, Fill Cliff's d from v25."
                break
            for r in runs12[1:]:
                new_p12.remove(r)
        ref12.addnext(new_p12)
        break

# Update "Fig 3.2" to "Fig 3.2 / 3.2b"
for s in slide12.shapes:
    if s.has_text_frame and s.text.strip() == "Fig 3.2":
        for para in s.text_frame.paragraphs:
            for run in para.runs:
                if "3.2" in run.text:
                    run.text = "Fig 3.2 / 3.2b"
        break

print("✓ Slide 12: Icon-to-T2 line + Fig 3.2b added")

# ── SLIDE 13 (index 12): Summary Table ──

slide13 = prs.slides[12]
tbl13 = find_table(slide13)
clone_last_row(tbl13, [
    "Icon to T2 (exposure learning)",
    "Fill from v25",
    "Fill",
])

print("✓ Slide 13: Q_IT2 summary row added")

# ── SAVE ──

prs.save(OUTPUT)
print(f"\nSaved to {OUTPUT}")
